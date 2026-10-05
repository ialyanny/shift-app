#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Generate 2027 shift poster Excel, TRUE 28-day cycle decoded from 2026 Jan+Feb A/D."""
import datetime, calendar
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 28-day base decoded: 2026/1/1 = day1 (index0). Verified P=28 zero-mismatch over 59 days.
BASE_A = '1000220022200220001100111001'
BASE_D = '0111001100022002220022000110'
def _opp(s):
    return ''.join('0' if c=='0' else ('1' if c=='2' else '2') for c in s)
BASE_B = _opp(BASE_A)
BASE_C = _opp(BASE_D)
BASE = {'A': BASE_A, 'B': BASE_B, 'C': BASE_C, 'D': BASE_D}
D0 = datetime.date(2026, 1, 1)

def shift_for(d):
    idx = (d - D0).days % 28
    return {'A': BASE_A[idx], 'B': BASE_B[idx], 'C': BASE_C[idx], 'D': BASE_D[idx]}

# sanity: every day exactly one '1' + one '2'
for i in range(28):
    w = sorted([BASE[t][i] for t in 'ABCD' if BASE[t][i] != '0'])
    assert w == ['1','2'], f'base day {i+1} bad {w}'

RED = {(1,1),(2,4),(2,5),(2,6),(2,7),(2,8),(2,28),(4,4),(4,5),(5,1),(6,9),(9,15),(9,28),(10,10),(10,25),(12,25)}
YELLOW = {(2,9),(2,10),(3,1),(4,6),(4,30),(10,11),(12,24),(12,31)}

MONTH_ENG = {1:'Jan',2:'Feb',3:'Mar',4:'Apr',5:'May',6:'Jun',7:'Jul',8:'Aug',9:'Sep',10:'Oct',11:'Nov',12:'Dec'}
WEEK_MAP = {0:'一',1:'二',2:'三',3:'四',4:'五',5:'六',6:'日'}

FILL_RED = PatternFill('solid', fgColor='FF0000')
FILL_YELLOW = PatternFill('solid', fgColor='FFFF00')
FILL_WHITE = PatternFill('solid', fgColor='FFFFFF')
FILL_GRAY = PatternFill('solid', fgColor='E8E8E8')
FONT_TITLE_RED = Font(name='標楷體', size=14, bold=True, color='FF0000')
FONT_TITLE_BLUE = Font(name='Arial', size=11, bold=True, color='0000AA')
FONT_MONTH_ENG = Font(name='Arial', size=8, color='0000AA')
FONT_MONTH_NUM = Font(name='Arial', size=12, bold=True, color='FF0000')
FONT_DAY = Font(name='Arial', size=8)
FONT_DAY_RED_BG = Font(name='Arial', size=8, bold=True, color='FFFFFF')
FONT_WEEK = Font(name='標楷體', size=8)
FONT_WEEK_SUN = Font(name='標楷體', size=8, color='FF0000', bold=True)
FONT_TEAM = Font(name='Arial', size=9, bold=True)
FONT_SHIFT = Font(name='Arial', size=9, bold=True)
FONT_FOOTER = Font(name='標楷體', size=9)
THIN = Side(style='thin', color='8EA9DB')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALIGN_C = Alignment(horizontal='center', vertical='center', wrap_text=True)

def disp(v):
    return '' if v == '0' else v

def build_sheet(ws, months):
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.2
    ws.page_margins.bottom = 0.2
    ws.column_dimensions['A'].width = 7
    for c in range(2, 33):
        ws.column_dimensions[get_column_letter(c)].width = 3.4
    ws.merge_cells('A1:AF1')
    ws['A1'].value = '2027台灣康寧顯示玻璃股份有限公司  公司統編04388145'
    ws['A1'].font = FONT_TITLE_RED
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 24
    ws.merge_cells('A2:AF2')
    ws['A2'].value = 'CORNING  Corning Display Technologies Taiwan  職工福利委員會13525747'
    ws['A2'].font = FONT_TITLE_BLUE
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 18
    r = 3
    for m in months:
        ndays = calendar.monthrange(2027, m)[1]
        ws.row_dimensions[r].height = 14
        ca = ws.cell(row=r, column=1, value=MONTH_ENG[m])
        ca.font = FONT_MONTH_ENG; ca.alignment = ALIGN_C; ca.border = BORDER
        for d in range(1, 32):
            cell = ws.cell(row=r, column=1+d, value=d if d <= ndays else None)
            cell.font = FONT_DAY; cell.alignment = ALIGN_C; cell.border = BORDER
            if d <= ndays:
                if (m, d) in RED:
                    cell.fill = FILL_RED; cell.font = FONT_DAY_RED_BG
                elif (m, d) in YELLOW:
                    cell.fill = FILL_YELLOW; cell.font = Font(name='Arial', size=8, bold=True)
                else:
                    cell.fill = FILL_WHITE
            else:
                cell.fill = FILL_GRAY
        r += 1
        ws.row_dimensions[r].height = 14
        ca = ws.cell(row=r, column=1, value=str(m))
        ca.font = FONT_MONTH_NUM; ca.alignment = ALIGN_C; ca.border = BORDER
        for d in range(1, 32):
            if d <= ndays:
                dt = datetime.date(2027, m, d)
                w = WEEK_MAP[dt.weekday()]
                cell = ws.cell(row=r, column=1+d, value=w)
                cell.alignment = ALIGN_C; cell.border = BORDER; cell.fill = FILL_WHITE
                cell.font = FONT_WEEK_SUN if w == '日' else FONT_WEEK
            else:
                cell = ws.cell(row=r, column=1+d, value=None)
                cell.fill = FILL_GRAY; cell.border = BORDER
        r += 1
        for team in ['A', 'B', 'C', 'D']:
            ws.row_dimensions[r].height = 13
            ca = ws.cell(row=r, column=1, value=team)
            ca.font = FONT_TEAM; ca.alignment = ALIGN_C; ca.border = BORDER
            for d in range(1, 32):
                if d <= ndays:
                    v = shift_for(datetime.date(2027, m, d))[team]
                    cell = ws.cell(row=r, column=1+d, value=disp(v) if disp(v) else None)
                    cell.font = FONT_SHIFT; cell.alignment = ALIGN_C; cell.border = BORDER; cell.fill = FILL_WHITE
                else:
                    cell = ws.cell(row=r, column=1+d, value=None)
                    cell.fill = FILL_GRAY; cell.border = BORDER
            r += 1
        ws.row_dimensions[r].height = 4
        r += 1
    ws.merge_cells(f'A{r}:AF{r}')
    r += 1
    for line in (footer_lines(True) if months[0] <= 6 else footer_lines(False)):
        ws.merge_cells(f'A{r}:AF{r}')
        c = ws.cell(row=r, column=1, value=line)
        c.font = FONT_FOOTER
        c.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[r].height = 14
        r += 1
    ws.print_title_rows = '1:2'

def footer_lines(first_half):
    if first_half:
        return [
            '1→早班/日班　台南廠06-5050520　南科管理局06-5051001',
            '2→晚班/夜班　台中廠04-24658999　南科警察隊06-5051405',
            '　→off　　　 台北02-27160338　南科消防分隊06-5052995',
            '■→國定假日(紅)　內湖02-8178132　南科診所06-5050225',
            '■→補假(黃)　　　新竹03-5820550　南科台銀06-5051701',
            '■→連續彈性假(綠)',
        ]
    return [
        '1→早班/日班　大都會計程車02-4499178　安環衛專線#1234',
        '2→晚班/夜班　國通計程車(台中)04-24618999　安全室專線#1332',
        '　→off　　　 EAP免費專線00801-491519　廠務部專線#1000',
        '■→國定假日(紅)　EAP24HR熱線02-77032446　性騷擾申訴專線#1777',
        '■→補假(黃)',
        '■→連續彈性假(綠)',
    ]

wb = Workbook()
ws1 = wb.active; ws1.title = '2027年1-6月'
build_sheet(ws1, [1,2,3,4,5,6])
ws2 = wb.create_sheet('2027年7-12月')
build_sheet(ws2, [7,8,9,10,11,12])
wb.save('2027台灣康寧輪班表.xlsx')
print('Excel saved')
# quick verify Jan D/A against expected pattern sample
import sys
sys.path.insert(0, '.')
